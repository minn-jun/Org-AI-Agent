"""진행 중 과제 문서를 STM · MTM · LTM 시드로 변환한다.

왜 필요한가
-----------
지금까지 STM · MTM은 **전부 합성**이었다. 종료 과제(20200504)에는 현재 시점
자료가 없어서다. 그래서 "3계층이 답변을 좋게 만든다"는 이 프로젝트의 핵심
주장에 실측 근거가 한 줄도 없었다.

이 과제는 2026-04-01 ~ 2027-09-30으로 **진행 중**이고, 세 계층에 해당하는
자료가 모두 실물로 있다.

  LTM  확정된 계획 · 협약 · 조직 공통 매뉴얼        doc_rag 패키지
  MTM  주차별 진행 산출물(주간회의 자료)            doc_rag 패키지 + seminar/
  STM  세미나 피드백 기록 · 킥오프 회의             회의내용/ 폴더만

길 A(계층별 청크 코퍼스)가 아니라 **길 B**를 쓴다 — 문서 하나를 md/json 한 장으로
내려 기존 시드와 같은 모양으로 만든다. 코드 수정이 0이고, "실제 MTM이 들어오면
무엇이 달라지나"를 먼저 볼 수 있다. 계층마다 다른 검색기를 붙이면 09-16에 겪은
계층 간 점수 눈금 문제가 되돌아온다.

실행
----
    python scripts/build_prentice_seed.py              # memory_seed_prentice/ 생성
    python scripts/build_prentice_seed.py --dry        # 배정만 출력
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sqlite3
import sys
import unicodedata
from collections import defaultdict
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]          # org_agent_mvp/
WORKSPACE = ROOT.parent                              # 중기청_조직지식AI플랫폼/
PACKAGE = WORKSPACE / "doc_rag_nochat_2026-09-22"
SQLITE = PACKAGE / "data" / "chunks.sqlite3"
FEEDBACK_DIR = WORKSPACE / "회의내용"
SEMINAR_DIR = WORKSPACE / "seminar"
OUT = ROOT / "memory_seed_prentice"

#: 시드 frontmatter의 `project` 값. 분석기의 필터 어휘가 여기서 나오므로
#: 실제 과제명과 맞아야 한다. 기관명이라 환경변수로 받는다.
PROJECT = os.environ.get("PRENTICE_PROJECT", "진행 과제")

# ─────────────────────────────────────────────────────────── 계층 배정 규칙
#
# 경로 조각으로만 판단한다. 파일명으로 판단하면 "작성중"이 협약 폴더에도 있어서
# 갈리지 않는다(협약 폴더의 `_archive`에 `작성중` 표기가 섞여 있다).

#: 색인에 넣지 않는다. 실명 · 계좌번호 · 급여액이 마스킹 없이 본문에 있다.
#: 선배 확인 전까지는 **빼는 쪽이 기본**이다.
EXCLUDE = (
    "연구비지출/01.영수증",
    "사업자등록증 및 통장 사본",
)
#: 파일명으로 걸러야 하는 개인정보. `05-협약/_archive/직원 연봉 계산.xlsx`에
#: 실명과 연봉액이 그대로 있었다. 폴더 규칙만으로는
#: 새는 자리가 있다.
PERSONAL_PATTERNS = ("연봉", "급여", "인건비", "명세서", "급여대장", "이체확인증")
#: 진행 중 산출물. 주차가 곧 판본이다.
MTM_PATHS = ("06-수행/주간 회의",)
#: 제출 · 승인된 공식 문서와 조직 공통 기준.
LTM_PATHS = (
    "03-제출",
    "05-협약",
    "연구비지출/00.자료",          # RCMS · 연구노트 작성 매뉴얼
    "01-서류/국가연구개발사업 협약서",
)

#: 서명용 서식과 증빙은 LTM에서 뺀다.
#:
#: 우리 LTM 정의는 "확정 · 공식 문서"다. 그런데 동의서 · 서약서 · 확인서처럼
#: **서명만 받는 서식**은 지식이 아니고, 기관별 사본이 3~4장씩 있어 검색 결과를
#: 같은 계열로 채운다. 09-16에 "정제되지 않은 초안과 빈 양식까지 LTM에 들어가
#: 있다"고 지적된 것이 이 부분이다.
#:
#: 빈칸 비율로 가려 보려 했으나 갈리지 않았다 — `연구개발계획서 요약문`(실내용)이
#: 0.70, `연구개발계획서 본문1`(빈 양식)이 0.67이었다. 표가 많은 문서는 둘 다
#: 빈칸이 많다. 그래서 **파일명 기준 명시 규칙**으로 두고, 이 목록 자체를
#: 사람이 검토하게 한다.
FORM_PATTERNS = (
    "동의서",
    "서약서",
    "확인서",
    "신청서",
    "체크리스트",
    "사업자등록증",
    "통장",
    "직인",
)
#: 내용을 확인한 빈 양식. 표 칸이 모두 비어 있다.
BLANK_FORMS = ("연구개발계획서 본문1", "검증 양식")

#: 회의 녹취(.vtt)는 **넣지 않는다.** (2026-10-07)
#:
#: 한 파일이 avgdl의 8배라 STM 검색을 거의 다 차지한다. 질문어가 대부분
#: 들어 있어 tf가 압도적이고, 짧은 피드백 문서가 경쟁에서 밀린다. 길이 보정을
#: 끝까지 올려도(b=1.0) 1위를 내주지 않았고, b=1.0은 STM tier@1을 12/15에서
#: 11/15로 떨어뜨렸다.
#:
#: 자료 자체도 음성 인식 잡음이 있다. "이종원"이 206회, 잘못 인식된 표기가
#: 1회 섞여 있고, 모델이 그 1회를 그대로 답에 옮긴 사례가 나왔다. 원문에
#: 있으니 할루시네이션은 아니지만, 사람 이름을 틀리게 답하는 것은 같다.
#:
#: 청크로 갈라 넣으면 길이 문제는 풀리지만 시드는 "파일 하나가 카드 하나"
#: 구조다. 구조를 바꾸기 전까지는 제외한다.
TRANSCRIPT_SUFFIX = ".vtt"

#: 특정 작성자의 발표자료는 PDF 추출에서 **띄어쓰기가 사라졌다**("컨텍스트구성").
#: 형태소 분석이 깨지므로 seminar/의 pptx 원본에서 다시 뽑는다.
#:
#: 대상 작성자와 저자 어휘는 실제 사람 이름이라 **코드에 두지 않는다.**
#: 쉼표로 구분해 환경변수로 넘긴다. 비우면 대체와 저자 표기를 건너뛴다.
#:     PRENTICE_PDF_AUTHOR="홍길동" PRENTICE_AUTHORS="홍길동,김철수"
BROKEN_PDF_AUTHOR = os.environ.get("PRENTICE_PDF_AUTHOR", "").strip()
KNOWN_AUTHORS = tuple(
    name.strip()
    for name in os.environ.get("PRENTICE_AUTHORS", "").split(",")
    if name.strip()
)

DATE_RE = re.compile(r"(20\d{2})[-_]?(\d{2})[-_]?(\d{2})")


def date_from_mtime(mtime: float | str | None) -> str | None:
    """패키지의 `mtime`은 유닉스 초다. 파일명에 날짜가 없을 때만 쓴다."""
    try:
        value = float(mtime)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return None
    if value <= 0:
        return None
    return datetime.fromtimestamp(value).strftime("%Y-%m-%d")


def iso_date(text: str) -> str | None:
    match = DATE_RE.search(text)
    if not match:
        return None
    year, month, day = match.groups()
    try:
        datetime(int(year), int(month), int(day))
    except ValueError:
        return None
    return f"{year}-{month}-{day}"


def slug(text: str) -> str:
    text = unicodedata.normalize("NFC", text)
    text = re.sub(r"[^0-9A-Za-z가-힣]+", "-", text).strip("-")
    return text[:60] or "doc"


def tier_of(rel_path: str) -> str | None:
    """경로로 계층을 정한다. None이면 넣지 않는다."""
    path = rel_path.replace("\\", "/")
    name = Path(path).name
    if any(part in path for part in EXCLUDE):
        return None
    if any(part in name for part in PERSONAL_PATTERNS):
        return None
    if any(part in name for part in FORM_PATTERNS):
        return None
    if any(part in name for part in BLANK_FORMS):
        return None
    # 녹취록은 계층을 주지 않는다. 위 주석 참고.
    if path.lower().endswith(TRANSCRIPT_SUFFIX):
        return None
    if any(part in path for part in MTM_PATHS):
        return "mtm"
    if any(part in path for part in LTM_PATHS):
        return "ltm"
    return None


def normalize_stem(file_name: str) -> str:
    """판본 중복을 묶는 열쇠.

    같은 문서의 hwp · hwpx · pdf · md가 나란히 있다. 전부 넣으면 검색 결과가
    같은 계열로 채워진다(실코퍼스에서 20.1% → 0%로 고친 문제다).
    """
    stem = Path(file_name).stem
    stem = re.sub(r"\s*\(?폰트깨짐방지\)?$", "", stem)
    # 판본 접미어를 **반복해서** 떼어 낸다.
    #   협약용_연구개발계획서_..._jic 수정_v2 → ..._jic 수정 → ...
    #   예산서_최종 협약본 · _작성중 2 · _협약용 복사본 → 예산서
    version = re.compile(
        r"[_\s-]*(?:v\d+|\d+차|jic\s*수정|수정|최종본?|작성중|복사본|편집완료"
        r"|평가요소편집완료|협약본|협약용|서명용|제출용|final)\s*\d*$",
        re.IGNORECASE,
    )
    while True:
        trimmed = version.sub("", stem).strip()
        if trimmed == stem or not trimmed:
            break
        stem = trimmed
    stem = re.sub(r"\s*\d+$", "", stem)
    return unicodedata.normalize("NFC", stem).strip().lower()


# ──────────────────────────────────────────────────────────── 원본 읽기
def load_package() -> list[dict]:
    if not SQLITE.exists():
        sys.exit(f"패키지를 찾지 못했다: {SQLITE}")
    conn = sqlite3.connect(SQLITE)
    conn.row_factory = sqlite3.Row
    # 파일명·경로의 정규화 형식이 섞여 있다(맥에서 만든 파일은 NFD).
    # `직원 연봉 계산.xlsx`은 NFD라 NFC 리터럴로는 걸리지 않았다.
    # 입구에서 전부 NFC로 맞춘다 — 규칙이 새지 않게.
    def nfc(value):
        return unicodedata.normalize("NFC", value) if isinstance(value, str) else value

    docs = []
    for row in conn.execute(
        "select doc_id, rel_path, file_name, format, n_chunks, mtime"
        "  from documents where n_chunks > 0"
    ):
        chunks = [
            {key: nfc(value) for key, value in dict(chunk).items()}
            for chunk in conn.execute(
                "select ord, text, headings, page_no from chunks"
                " where doc_id = ? order by ord",
                (row["doc_id"],),
            )
        ]
        docs.append(
            {**{key: nfc(value) for key, value in dict(row).items()}, "chunks": chunks}
        )
    conn.close()
    return docs


def pptx_text(path: Path) -> str:
    """슬라이드 순서대로 글을 뽑는다. 표도 포함한다."""
    from pptx import Presentation

    parts: list[str] = []
    for index, slide in enumerate(Presentation(path).slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        lines.append(text)
            if getattr(shape, "has_table", False):
                for trow in shape.table.rows:
                    cells = [cell.text.strip() for cell in trow.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
        if lines:
            parts.append(f"## 슬라이드 {index}\n\n" + "\n".join(lines))
    return "\n\n".join(parts)


def body_from_chunks(chunks: list[dict]) -> str:
    """청크를 ord 순서로 이어 붙인다. headings가 바뀌면 소제목으로 남긴다."""
    parts: list[str] = []
    last_heading = None
    for chunk in chunks:
        heading = (chunk.get("headings") or "").strip()
        if heading and heading != last_heading:
            parts.append(f"## {heading}")
            last_heading = heading
        parts.append(str(chunk.get("text") or "").strip())
    return "\n\n".join(part for part in parts if part)


def first_sentence(body: str, limit: int = 160) -> str:
    text = re.sub(r"^#.*$", "", body, flags=re.MULTILINE)
    text = re.sub(r"\s+", " ", text).strip()
    return text[:limit]


def write_md(path: Path, meta: dict, body: str) -> None:
    lines = ["---"]
    for key, value in meta.items():
        lines.append(f"{key}: {value}")
    lines.append("---")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n\n" + body.strip() + "\n", encoding="utf-8")


# ──────────────────────────────────────────────────────────── 변환
def convert_package(dry: bool) -> dict:
    docs = load_package()
    assigned: dict[str, list[dict]] = defaultdict(list)
    skipped: list[tuple[str, str]] = []

    for doc in docs:
        tier = tier_of(doc["rel_path"])
        if tier is None:
            skipped.append((doc["file_name"], "경로 규칙에서 제외"))
            continue
        assigned[tier].append(doc)

    # 판본 중복을 계층 안에서 묶는다. 청크가 가장 많은 쪽을 남긴다.
    kept: dict[str, list[dict]] = {}
    folded: list[tuple[str, str]] = []
    for tier, items in assigned.items():
        groups: dict[str, list[dict]] = defaultdict(list)
        for doc in items:
            groups[normalize_stem(doc["file_name"])].append(doc)
        chosen = []
        for group in groups.values():
            group.sort(key=lambda d: (-d["n_chunks"], d["file_name"]))
            chosen.append(group[0])
            for other in group[1:]:
                folded.append((other["file_name"], f"→ {group[0]['file_name']}"))
        kept[tier] = sorted(chosen, key=lambda d: d["file_name"])

    written: dict[str, list[str]] = defaultdict(list)
    replaced: list[str] = []

    for tier, items in kept.items():
        for doc in items:
            date = iso_date(doc["file_name"]) or date_from_mtime(doc["mtime"]) or ""
            body = body_from_chunks(doc["chunks"])
            source_note = f"doc_rag 패키지 / {doc['rel_path']}"

            # 지정한 작성자의 발표자료는 pptx 원본에서 다시 뽑는다
            if (
                tier == "mtm"
                and BROKEN_PDF_AUTHOR
                and BROKEN_PDF_AUTHOR in doc["file_name"]
                and date
            ):
                candidates = sorted(SEMINAR_DIR.glob(f"{date}*.pptx"))
                if candidates:
                    text = pptx_text(candidates[0])
                    if len(text) > 200:
                        body = text
                        source_note = f"seminar/{candidates[0].name} (PDF 띄어쓰기 손실 대체)"
                        replaced.append(f"{date} {candidates[0].name}")

            # 녹취록을 빼면서 이 경로로 들어오는 stm은 없어졌다. STM은
            # 회의내용/ 폴더에서만 들어온다(킥오프 1 · 세미나 피드백 10).
            if tier == "mtm":
                meta_type = "weekly_meeting"
            elif "매뉴얼" in doc["file_name"]:
                meta_type = "manual"
            elif "계획서" in doc["file_name"]:
                meta_type = "research_plan"
            else:
                meta_type = "official_document"

            author = next(
                (name for name in KNOWN_AUTHORS if name in doc["file_name"]), ""
            )
            title = Path(doc["file_name"]).stem
            meta = {
                "memory_tier": tier,
                "source_type": meta_type,
                "project": PROJECT,
                "title": title,
                "date": date,
                "permission_scope": "project_team",
                "status": "active",
                "summary": first_sentence(body),
                "summary_source": "auto",
                "origin": source_note,
            }
            if author:
                meta["author"] = author
            name = f"{date or 'undated'}-{slug(title)}.md"
            written[tier].append(name)
            if not dry:
                write_md(OUT / tier / name, meta, body)

    return {
        "written": written,
        "skipped": skipped,
        "folded": folded,
        "replaced": replaced,
    }


def convert_feedback(dry: bool) -> list[str]:
    """회의내용/*.txt → STM. 실제로 받은 피드백 기록이다."""
    if not FEEDBACK_DIR.exists():
        return []
    names = []
    seen_digest: dict[str, str] = {}
    for path in sorted(FEEDBACK_DIR.glob("*.txt")):
        date = iso_date(path.name)
        text = path.read_text(encoding="utf-8").replace("\r\n", "\n").strip()
        digest = str(hash(text))
        duplicate_of = seen_digest.get(digest)
        seen_digest.setdefault(digest, path.name)

        kind = "kickoff_note" if "킥오프" in path.name else "seminar_feedback"
        title = f"{date} {'킥오프 회의' if kind == 'kickoff_note' else '세미나 피드백'}"
        payload = {
            "memory_tier": "stm",
            "source_type": kind,
            "project": PROJECT,
            "title": title,
            "date": date or "",
            "permission_scope": "project_team",
            "status": "active",
            "summary": first_sentence(text),
            "summary_source": "auto",
            "origin": f"회의내용/{path.name}",
            "body": text,
        }
        if duplicate_of:
            # 09-18 파일이 09-03의 사본이다(md5 동일). 내용이 비어 있는 셈이라
            # 표시해 둔다 — 모르고 쓰면 그 주차 피드백이 있다고 착각한다.
            payload["status"] = "suspect_duplicate"
            payload["duplicate_of"] = duplicate_of
        label = "킥오프-회의" if kind == "kickoff_note" else "세미나-피드백"
        name = f"{date or 'undated'}-{label}.json"
        names.append(name + ("  (사본 의심)" if duplicate_of else ""))
        if not dry:
            target = OUT / "stm" / name
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(
                json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
            )
    return names


def extra_seminar_weeks(dry: bool, existing: set[str]) -> list[str]:
    """패키지 스냅숏(09-22) 이후·누락 주차를 seminar/에서 보탠다."""
    names = []
    for path in sorted(SEMINAR_DIR.glob("*.pptx")):
        date = iso_date(path.name)
        if not date or date in existing:
            continue
        body = pptx_text(path)
        if len(body) < 200:
            continue
        title = path.stem
        meta = {
            "memory_tier": "mtm",
            "source_type": "weekly_meeting",
            "project": PROJECT,
            "title": title,
            "date": date,
            "permission_scope": "project_team",
            "status": "active",
            "summary": first_sentence(body),
            "summary_source": "auto",
            "origin": f"seminar/{path.name} (패키지 색인에 없음)",
            "author": BROKEN_PDF_AUTHOR,
        }
        name = f"{date}-{slug(title)}.md"
        names.append(name)
        if not dry:
            write_md(OUT / "mtm" / name, meta, body)
    return names


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dry", action="store_true", help="배정만 출력하고 쓰지 않는다")
    args = parser.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    if not args.dry and OUT.exists():
        for old in OUT.rglob("*"):
            if old.is_file():
                old.unlink()

    result = convert_package(args.dry)
    written = result["written"]
    # 파일명이 날짜로 시작한다. seminar에서 보탤 때 중복을 피하려고 모아 둔다.
    dates = {
        name[:10]
        for names in written.values()
        for name in names
        if DATE_RE.match(name)
    }
    extra = extra_seminar_weeks(args.dry, dates)
    written["mtm"].extend(extra)
    feedback = convert_feedback(args.dry)
    written["stm"].extend(feedback)

    print(f"출력: {OUT}{'  (--dry, 쓰지 않음)' if args.dry else ''}")
    for tier in ("stm", "mtm", "ltm"):
        print(f"\n── {tier.upper()}  {len(written[tier])}건")
        for name in sorted(written[tier]):
            print(f"   {name}")
    print(f"\n── seminar에서 보탠 주차 {len(extra)}건")
    for name in extra:
        print(f"   {name}")
    print(f"\n── PDF를 pptx 원본으로 대체 {len(result['replaced'])}건")
    for name in result["replaced"]:
        print(f"   {name}")
    print(f"\n── 판본 중복으로 접음 {len(result['folded'])}건")
    for name, into in result["folded"][:40]:
        print(f"   {name[:58]:<60}{into[:40]}")
    print(f"\n── 제외 {len(result['skipped'])}건 (개인정보 · 규칙 밖)")
    for name, why in result["skipped"][:40]:
        print(f"   {name[:58]:<60}{why}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
